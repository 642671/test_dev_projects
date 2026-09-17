import json
import os
from pathlib import Path

from openpyxl import load_workbook

HEADERS = [
    "编号",
    "模块",
    "用例名称",
    "优先级",
    "前置条件",
    "操作步骤",
    "输入数据",
    "预期结果",
    "验证结果",
    "备注",
    "bug",
]

LIST_FIELDS = {"前置条件", "操作步骤", "输入数据", "预期结果"}
FAILURE_NOTICE = (
    "快照任务删除失败！任务中包含受保护的快照，"
    "请前往快照清单进行解锁或等待保护期结束后再重试。"
)

OVERRIDES = {
    "SNAP-IMM-026": {
        "用例名称": "删除快照-任务仅含普通快照",
        "前置条件": [
            "1. 已使用管理员账号登录TOS，快照任务页面可正常打开；",
            "2. 已存在仅包含普通快照的快照任务1个。",
        ],
        "操作步骤": [
            "1. 进入快照-快照任务页面；",
            "2. 选中仅包含普通快照的任务并点击删除；",
            "3. 在确认弹窗中点击确认。",
        ],
        "输入数据": [
            "1. 快照任务：snap-task-026（示例值）；",
            "2. 任务内快照：2条，均未锁定且未启用不可变保护；",
            "3. 操作：删除任务并确认。",
        ],
        "预期结果": [
            "1. 页面弹出“操作成功！”提示；",
            "2. 快照任务删除成功，列表刷新后该任务不再显示；",
            "3. 重新进入页面后任务仍不存在，任务内未受保护的普通快照已按规则清理。",
        ],
        "验证结果": "未执行",
        "备注": "snap-task-026为示例值，执行时可根据实际测试环境替换",
        "bug": "",
    },
    "SNAP-IMM-027": {
        "用例名称": "删除快照-任务含锁定快照但不含不可变快照",
        "前置条件": [
            "1. 已使用管理员账号登录TOS，快照任务页面可正常打开；",
            "2. 已存在包含锁定快照且不含不可变快照的快照任务1个。",
        ],
        "操作步骤": [
            "1. 进入快照-快照任务页面；",
            "2. 选中包含锁定快照且不含不可变快照的任务并点击删除；",
            "3. 在确认弹窗中点击确认。",
        ],
        "输入数据": [
            "1. 快照任务：snap-task-027（示例值）；",
            "2. 任务内快照：2条，其中至少1条已锁定，均未启用不可变保护；",
            "3. 操作：删除任务并确认。",
        ],
        "预期结果": [
            f"1. 页面弹出“{FAILURE_NOTICE}”提示；",
            "2. 快照任务删除失败，任务及任务内快照均保留；",
            "3. 刷新后任务和锁定状态不变，不发生部分删除。",
        ],
        "验证结果": "未执行",
        "备注": "snap-task-027为示例值，执行时可根据实际测试环境替换",
        "bug": "",
    },
    "SNAP-IMM-028": {
        "用例名称": "删除快照-任务含不可变快照但不含锁定快照",
        "前置条件": [
            "1. 已使用管理员账号登录TOS，快照任务页面可正常打开；",
            "2. 已存在包含保护期内不可变快照且不含锁定快照的快照任务1个。",
        ],
        "操作步骤": [
            "1. 进入快照-快照任务页面；",
            "2. 选中包含保护期内不可变快照且不含锁定快照的任务并点击删除；",
            "3. 在确认弹窗中点击确认。",
        ],
        "输入数据": [
            "1. 快照任务：snap-task-028（示例值）；",
            "2. 任务内快照：2条，其中至少1条不可变且保护期未结束，均未锁定；",
            "3. 操作：删除任务并确认。",
        ],
        "预期结果": [
            f"1. 页面弹出“{FAILURE_NOTICE}”提示；",
            "2. 快照任务删除失败，任务及任务内快照均保留；",
            "3. 刷新后任务和不可变保护状态不变，不发生部分删除。",
        ],
        "验证结果": "未执行",
        "备注": "snap-task-028为示例值，执行时可根据实际测试环境替换",
        "bug": "",
    },
    "SNAP-IMM-029": {
        "用例名称": "删除快照-任务同时含锁定与不可变快照",
        "前置条件": [
            "1. 已使用管理员账号登录TOS，快照任务页面可正常打开；",
            "2. 已存在同时包含锁定快照和保护期内不可变快照的快照任务1个。",
        ],
        "操作步骤": [
            "1. 进入快照-快照任务页面；",
            "2. 选中同时包含锁定快照和保护期内不可变快照的任务并点击删除；",
            "3. 在确认弹窗中点击确认。",
        ],
        "输入数据": [
            "1. 快照任务：snap-task-029（示例值）；",
            "2. 任务内快照：3条，其中至少1条已锁定、至少1条不可变且保护期未结束；",
            "3. 操作：删除任务并确认。",
        ],
        "预期结果": [
            f"1. 页面弹出“{FAILURE_NOTICE}”提示；",
            "2. 快照任务删除失败，任务及任务内快照均保留；",
            "3. 刷新后锁定与不可变保护状态均不变，不发生部分删除。",
        ],
        "验证结果": "未执行",
        "备注": "snap-task-029为示例值，执行时可根据实际测试环境替换",
        "bug": "",
    },
}


def normalize_cell(value, field):
    if value is None:
        return [] if field in LIST_FIELDS else ""
    text = str(value).strip()
    if field in LIST_FIELDS:
        return [line.strip() for line in text.splitlines() if line.strip()]
    return text


def main():
    workbook_path = Path(os.environ["WB_PATH"])
    cases_path = Path(os.environ["CASES_JSON"])
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    workbook.close()

    records = []
    for row in rows[1:]:
        values = list(row) + [None] * (len(HEADERS) - len(row))
        record = {
            header: normalize_cell(values[index], header)
            for index, header in enumerate(HEADERS)
        }
        override = OVERRIDES.get(record["编号"])
        if override:
            record.update(override)
        records.append(record)

    cases_path.write_text(
        json.dumps(records, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(f"updated {cases_path} with {len(records)} cases")


if __name__ == "__main__":
    main()
